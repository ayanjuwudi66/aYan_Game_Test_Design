#!/usr/bin/env python3
"""Generate a game-test-design testcase workbook.

两种格式：
  1) 内部骨架（默认，12 列）：用例ID|模块/功能|状态节点|前置条件|操作步骤|预期结果|
     测试类型|打断层级|影响面|风险级别|优先级|备注。供推演阶段使用。
  2) 交付版（--format delivery，8 列）：功能|测试点|前提条件|操作步骤|预期结果|备注|优先级|结果。
     功能列按连续相同值合并；备注/结果留空。这是 SKILL.md「输出约定」规定的用户交付格式，
     可选附带「待确认项清单」「图源说明」两个 sheet。

Usage:
  python gen_testcase_xlsx.py -o cases.xlsx                     # 12 列内部骨架
  python gen_testcase_xlsx.py -i cases.json -o cases.xlsx        # 12 列，读 JSON
  python gen_testcase_xlsx.py -i delivery.json -o out.xlsx --format delivery   # 8 列交付版

内部骨架 JSON：对象数组，键为 12 列名（未知键忽略、缺列留空）。

交付版 JSON（两种形态）：
  A) 纯用例：对象数组，键为 功能/测试点/前提条件/操作步骤/预期结果/优先级（优先级可省，默认 P2）。
     只产「测试用例」sheet；同「功能」连续多条自动合并功能列。
     「前提条件」可省略——若「操作步骤」以「前置：…」开头，脚本自动拆分为两列（见 split_precondition）。
  B) 三件套：{"cases": [...], "pending": [...], "img_notes": [...]}
     - cases：同上对象数组（必填）
     - pending（可选）：对象数组，键 = 编号/分类/待确认/来源/影响/默认处理/状态
     - img_notes（可选）：对象数组，键 = 标题/内容
     依次产「测试用例」「待确认项清单」「图源说明」三个 sheet。
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


COLUMNS = [
    "用例ID", "模块/功能", "状态节点", "前置条件", "操作步骤", "预期结果",
    "测试类型", "打断层级", "影响面", "风险级别", "优先级", "备注",
]
DELIVERY_COLUMNS = ["功能", "测试点", "前提条件", "操作步骤", "预期结果", "备注", "优先级", "结果"]
# 含多条时需要「单元格内换行」的列（见 cell_break）
BREAK_COLUMNS = {"前置条件", "前提条件", "操作步骤", "预期结果"}
LEVELS = ["无", "L1", "L2", "L3", "L4", "L5", "L6-a", "L6-b", "L7"]
# 单元格内换行切分点：序号标记（1. / 2、/ 3) / 4．/ 5））之前；序号后紧跟数字则不算序号
STEP_SPLIT_RE = re.compile(r'[ \t]*(?<!\d)(?<!第)(?=\d{1,2}[.、．)）](?!\d))')
IMPACTS = ["自身", "队友", "敌对玩家", "中立单位", "环境场景", "观战/旁观者", "服务器全局", "AI/召唤物"]


def load_rows(path: Path | None):
    if path is None:
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    return data


def build_workbook(rows):
    if not isinstance(rows, list):
        raise ValueError("skeleton 模式输入 JSON 必须是对象数组")
    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases"
    ws.append(COLUMNS)
    for idx, row in enumerate(rows, 1):
        values = [row.get(col, "") for col in COLUMNS]
        if not values[0]:
            values[0] = f"GT-{idx:03d}"
        values = [cell_break(v) if col in BREAK_COLUMNS else v
                  for col, v in zip(COLUMNS, values)]
        ws.append(values)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [12, 18, 18, 28, 36, 42, 16, 14, 20, 12, 10, 30]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Keep the framework vocabulary visible and reusable in the workbook.
    dictionary = wb.create_sheet("Dictionaries")
    dictionary.append(["打断层级", "影响面"])
    max_len = max(len(LEVELS), len(IMPACTS))
    for i in range(max_len):
        dictionary.append([
            LEVELS[i] if i < len(LEVELS) else "",
            IMPACTS[i] if i < len(IMPACTS) else "",
        ])
    for cell in dictionary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    dictionary.column_dimensions["A"].width = 14
    dictionary.column_dimensions["B"].width = 18

    wb.defined_names.add(
        DefinedName("InterruptionLevels", attr_text="'Dictionaries'!$A$2:$A$10")
    )
    level_validation = DataValidation(
        type="list", formula1="=InterruptionLevels", allow_blank=True
    )
    ws.add_data_validation(level_validation)
    level_validation.add("H2:H1048576")
    return wb


PENDING_COLUMNS = ["编号", "分类", "待确认", "来源", "影响", "默认处理", "状态"]
IMG_NOTE_COLUMNS = ["标题", "内容"]


def _style_header(ws, ncols):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1][:ncols]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def build_pending_sheet(wb, pending):
    """待确认项清单：编号|分类|待确认|来源|影响|默认处理|状态"""
    ws = wb.create_sheet("待确认项清单")
    ws.append(PENDING_COLUMNS)
    _style_header(ws, len(PENDING_COLUMNS))
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in pending:
        if isinstance(row, dict):
            values = [
                row.get("编号", ""), row.get("分类", ""), row.get("待确认", ""),
                row.get("来源", ""), row.get("影响", ""), row.get("默认处理", ""),
                row.get("状态", "待确认"),
            ]
        else:
            values = list(row)
        ws.append(values)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ws.max_row + 1):
        for c in (1, 2, 7):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, width in enumerate([8, 8, 62, 40, 40, 46, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def build_img_note_sheet(wb, img_notes):
    """图源说明：标题|内容（两栏式说明）"""
    ws = wb.create_sheet("图源说明")
    ws.append(IMG_NOTE_COLUMNS)
    _style_header(ws, len(IMG_NOTE_COLUMNS))
    for row in img_notes:
        if isinstance(row, dict):
            ws.append([row.get("标题", ""), row.get("内容", "")])
        else:
            ws.append(list(row))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 96
    return ws


def split_precondition(steps):
    """把「前置：…。1. …」拆成 (前提条件, 操作步骤)。

    历史格式把前置条件写进操作步骤首句（`前置：<条件>。1. …`），导致列职责混杂、
    执行人无法一眼区分"要准备什么"与"要做什么"。2026-09-03 起交付版拆为独立两列。

    规则：以「前置：」开头时，取到第一个「。」为前提条件，其后为操作步骤。
    不以「前置：」开头（或已显式提供「前提条件」）时原样返回。
    """
    if not isinstance(steps, str):
        return "", steps or ""
    s = steps.strip()
    for marker in ("前置：", "前置:"):
        if s.startswith(marker):
            body = s[len(marker):]
            idx = body.find("。")
            if idx == -1:
                return body.strip(), ""
            return body[:idx].strip(), body[idx + 1:].strip()
    return "", s


def cell_break(text):
    """把单元格内并列的多条内容拆成**单元格内换行**，避免「1. … 2. … 3. …」堆成一长行。

    需求（2026-09-03）：前提条件 / 操作步骤 / 预期结果 等列若含多条，必须在**同一个
    单元格内**逐条换行显示（1.\\n2.\\n3.），而不是挤在一行靠列宽折行——后者序号会被
    折断、执行人一眼看不出总共几步。

    规则：
      - 先按已有换行切分逐行处理，已手写换行的保持原样，不重复插入空行。
      - 在序号标记（1. / 2、/ 3) / 4．/ 5））前断行；序号后紧跟数字时**不算**序号，
        以此排除「3.5 秒」「40.50%」这类小数不被误切。
      - 「第2.」这类序数词不断行。
      - 只处理**带序号**的并列条目；用「；」分隔但无序号的文案保持原样，不臆测拆分。
    """
    if not isinstance(text, str):
        return text or ""
    out = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        for part in STEP_SPLIT_RE.sub("\n", line).split("\n"):
            part = part.strip()
            if part:
                out.append(part)
    return "\n".join(out)


def build_delivery_workbook(rows, sheet_name="测试用例"):
    """8 列交付版，可含「待确认项清单」「图源说明」两个附带 sheet。

    列：功能|测试点|前提条件|操作步骤|预期结果|备注|优先级|结果。

    rows 支持两种输入：
      - list[dict]：纯用例，只产用例 sheet（键：功能/测试点/前提条件/操作步骤/预期结果/优先级）
      - dict：{"cases": [...], "pending": [...], "img_notes": [...]}
              pending 与 img_notes 可选；pending 键 = 编号/分类/待确认/来源/影响/默认处理/状态，
              img_notes 键 = 标题/内容。
    sheet_name 为用例工作表名，默认「测试用例」，可按被测功能命名（如「采收功能测试用例」）。
    """
    cases = rows if isinstance(rows, list) else rows.get("cases", [])

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(sheet_name)
    ws.append(DELIVERY_COLUMNS)
    _style_header(ws, len(DELIVERY_COLUMNS))

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in cases:
        if not isinstance(row, dict):
            raise ValueError("交付版用例必须是对象数组")
        func = row.get("功能", "")
        point = row.get("测试点", "")
        steps = row.get("操作步骤", "")
        expect = row.get("预期结果", "")
        prio = row.get("优先级", "P2")
        # 前提条件独立成列；未显式提供时从操作步骤首句「前置：…」自动拆分
        pre = row.get("前提条件")
        if not pre:
            pre, steps = split_precondition(steps)
        # 前提条件/操作步骤/预期结果：含多条时在单元格内换行，不堆成一行
        ws.append([func, point, cell_break(pre), cell_break(steps), cell_break(expect),
                   "", prio, ""])

    # 功能列合并：连续相同「功能」合并单元格
    last_func = None
    run_start = 2
    merge_runs = []
    for r in range(2, ws.max_row + 1):
        func = ws.cell(row=r, column=1).value
        if func != last_func:
            if last_func is not None and r - 1 >= run_start:
                merge_runs.append((run_start, r - 1))
            run_start = r
            last_func = func
    if ws.max_row >= run_start:
        merge_runs.append((run_start, ws.max_row))
    for start, end in merge_runs:
        if end > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        c = ws.cell(row=start, column=1)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ws.max_row + 1):
        for c in (7, 8):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    for i, width in enumerate([24, 30, 40, 48, 52, 20, 8, 8], 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 附带 sheet
    if isinstance(rows, dict) and rows.get("pending"):
        build_pending_sheet(wb, rows["pending"])
    if isinstance(rows, dict) and rows.get("img_notes"):
        build_img_note_sheet(wb, rows["img_notes"])
    return wb


def main():
    parser = argparse.ArgumentParser(description="生成游戏测试用例 Excel（12 列骨架 / 8 列交付版）")
    parser.add_argument("-i", "--input", type=Path, help="用例 JSON 文件（对象数组或对象）")
    parser.add_argument("-o", "--output", type=Path, required=True, help="输出 .xlsx 路径")
    parser.add_argument("--format", choices=["skeleton", "delivery"], default="skeleton",
                        help="skeleton=12 列内部骨架（默认）；delivery=8 列交付版（可含待确认项/图源说明）")
    parser.add_argument("--sheet-name", type=str, default="测试用例",
                        help="delivery 模式用例工作表名（默认「测试用例」，可按被测功能命名）")
    args = parser.parse_args()
    rows = load_rows(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.format == "delivery":
        wb = build_delivery_workbook(rows, sheet_name=args.sheet_name)
        wb.save(args.output)
        n_case = len(rows) if isinstance(rows, list) else len(rows.get("cases", []))
        n_pend = 0 if isinstance(rows, list) else len(rows.get("pending") or [])
        n_img = 0 if isinstance(rows, list) else len(rows.get("img_notes") or [])
        print(f"已生成: {args.output}（{n_case} 条用例 + {n_pend} 项待确认 + {n_img} 行图源说明，format=delivery）")
    else:
        build_workbook(rows).save(args.output)
        print(f"已生成: {args.output}（{len(rows)} 条用例，format=skeleton）")


if __name__ == "__main__":
    main()
